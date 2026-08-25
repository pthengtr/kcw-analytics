"""Write the 10-sheet audit workbook. Never mutate input files."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .config import INSUFFICIENT_DATA, PROGRAM_NAME, PROGRAM_VERSION
from .reconciliation import ReconciliationResult

MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=16, color="1F4E79")
LABEL_FONT = Font(bold=True, name="Calibri", size=11)
NORMAL_FONT = Font(name="Calibri", size=11)
TOTAL_FILL = PatternFill("solid", fgColor="D6DCE4")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
PASS_FONT = Font(color="006100", bold=True, name="Calibri")
WARN_FONT = Font(color="9C5700", bold=True, name="Calibri")
FAIL_FONT = Font(color="9C0006", bold=True, name="Calibri")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    if value is INSUFFICIENT_DATA or value == INSUFFICIENT_DATA:
        return INSUFFICIENT_DATA
    if isinstance(value, float) and value != value:  # NaN
        return None
    try:
        import pandas as pd

        if pd.isna(value):
            return None
    except Exception:  # noqa: BLE001
        pass
    return value


def _is_money_column(name: str) -> bool:
    key = name.lower()
    tokens = (
        "amount",
        "price",
        "fee",
        "income",
        "refund",
        "adjustment",
        "difference",
        "gap",
        "balance",
        "withdrawal",
        "inflow",
        "ยอด",
    )
    return any(token in key for token in tokens)


def _write_header(ws, headers: list[str], row: int = 1) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def _autosize(ws, min_width: int = 12, max_width: int = 42) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        length = 0
        for cell in column[:80]:
            if cell.value is None:
                continue
            length = max(length, min(len(str(cell.value)), max_width))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 3))


def _apply_status_fill(cell, value: Any) -> None:
    text = str(value or "")
    if text in {"PASS", "MATCHED"}:
        cell.fill = PASS_FILL
        cell.font = PASS_FONT
    elif text in {"WARNING", "NET_SETTLED", "ORDER_NOT_RELEASED", "INFO"}:
        cell.fill = WARN_FILL
        cell.font = WARN_FONT
    elif text in {"FAIL", "AMOUNT_MISMATCH", "UNKNOWN"} or "FAIL" in text:
        cell.fill = FAIL_FILL
        cell.font = FAIL_FONT
    elif text in {"CANCELLED_OR_REFUNDED", "FINANCE_FROM_OTHER_PERIOD"}:
        cell.fill = WARN_FILL
        cell.font = WARN_FONT


def _write_table(ws, headers: list[str], records: list[dict[str, Any]], money_cols: set[str] | None = None) -> int:
    _write_header(ws, headers)
    money_cols = money_cols or {h for h in headers if _is_money_column(h)}
    for r_idx, record in enumerate(records, start=2):
        for c_idx, header in enumerate(headers, start=1):
            value = _excel_value(record.get(header))
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = NORMAL_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if header in money_cols and isinstance(value, (int, float)):
                cell.number_format = MONEY_FORMAT
            if header.lower() in {"status", "match_status", "validation_status", "severity"} or header == "status":
                _apply_status_fill(cell, value)
    last_row = 1 + len(records)
    if records:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    return last_row


def _add_generated_total(ws, headers: list[str], records: list[dict[str, Any]], money_cols: set[str], start_row: int) -> None:
    if not records:
        return
    total_row = start_row + 1
    for c_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=total_row, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.font = LABEL_FONT
        cell.border = THIN
        if c_idx == 1:
            cell.value = "GENERATED TOTAL"
            continue
        if header in money_cols:
            letter = get_column_letter(c_idx)
            cell.value = f"=SUM({letter}2:{letter}{start_row})"
            cell.number_format = MONEY_FORMAT


def _records_from_frame(frame: Any) -> list[dict[str, Any]]:
    if frame is None or getattr(frame, "empty", True):
        return []
    return frame.to_dict(orient="records")


def _write_kv(ws, start_row: int, items: list[tuple[str, Any]], money_labels: set[str] | None = None) -> int:
    money_labels = money_labels or set()
    row = start_row
    for label, value in items:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=1).alignment = Alignment(vertical="center")
        cell = ws.cell(row=row, column=2, value=_excel_value(value))
        cell.font = NORMAL_FONT
        if label in money_labels or _is_money_column(label):
            if isinstance(_excel_value(value), (int, float)):
                cell.number_format = MONEY_FORMAT
        if label.lower().endswith("สถานะ") or str(value) in {"PASS", "WARNING", "FAIL"}:
            _apply_status_fill(cell, value)
        row += 1
    return row


def write_report(result: ReconciliationResult, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    _write_executive(wb.active, result)
    _write_order_summary(wb.create_sheet("Order_Summary"), result)
    _write_finance_summary(wb.create_sheet("Finance_Summary"), result)
    _write_order_vs_finance(wb.create_sheet("Order_vs_Finance"), result)
    _write_fee_details(wb.create_sheet("Fee_Details"), result)
    _write_wallet(wb.create_sheet("Wallet_Summary"), result)
    _write_exceptions(wb.create_sheet("Exceptions"), result)
    _write_unknown(wb.create_sheet("Unknown_Mappings"), result)
    _write_validations(wb.create_sheet("Validations"), result)
    _write_methodology(wb.create_sheet("Methodology"), result)

    wb.save(output_path)
    return output_path


def _write_executive(ws, result: ReconciliationResult) -> None:
    ws.title = "Executive_Summary"
    ws["A1"] = "สรุปกระทบยอดรายได้และเงินรับ Lazada"
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.print_title_rows = "1:2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    banner = result.overall_status
    ws["A2"] = f"สถานะรวม: {banner}"
    _apply_status_fill(ws["A2"], banner)
    if banner == "FAIL":
        ws["B2"] = "สร้างรายงานแล้วแต่การตรวจสอบไม่ผ่าน — ห้ามรายงานว่ากระทบยอดสำเร็จ"
        ws["B2"].font = FAIL_FONT
    elif banner == "WARNING":
        ws["B2"] = "สร้างรายงานแล้ว มี WARNING ต้องตรวจก่อนลงบัญชี"
        ws["B2"].font = WARN_FONT
    else:
        ws["B2"] = "ผ่านการตรวจสอบเบื้องต้น — นักบัญชียังต้องตรวจผลอีกครั้ง (ไม่ใช่คำรับรองทางบัญชี)"
        ws["B2"].font = PASS_FONT

    expenses = result.expense_summary
    wallet = result.wallet_totals
    matched = result.match_counts.get("MATCHED", 0)
    exception_count = 0 if result.exceptions.empty else len(result.exceptions)
    money_labels = {
        "ยอดขายรวม (paidPrice)",
        "รายได้เข้า Wallet (finance_income)",
        "ค่าขนส่ง",
        "ค่าบริการ",
        "ค่าธรรมเนียมอื่น",
        "คืนเงิน",
        "รายการปรับปรุง",
        "รายการไม่ทราบประเภท",
        "ยอดสุทธิที่ควรเข้า Wallet",
        "ยอดโอนเข้าธนาคาร (ค่าบวกเพื่ออ่านง่าย)",
        "ยอดยกมา",
        "ยอดคงเหลือคำนวณ",
        "ผลต่าง Wallet",
        "ยอดถอน Auto Withdrawal (ค่าบวก)",
    }
    items = [
        ("เดือนที่ประมวลผล", result.month),
        ("เวลาประมวลผล", result.generated_at),
        ("เวอร์ชันโปรแกรม", f"{PROGRAM_NAME} {PROGRAM_VERSION}"),
        ("ไฟล์คำสั่งซื้อ", ", ".join(result.source_files.get("orders") or [])),
        ("ไฟล์รายการทางบัญชี", ", ".join(result.source_files.get("finance") or [])),
        ("ไฟล์ยอดของฉัน", ", ".join(result.source_files.get("wallet") or [])),
        ("โปรไฟล์ Finance", result.finance_profile or ""),
        ("จำนวนออเดอร์ (Order_Summary)", len(result.order_summary)),
        ("ยอดขายรวม (paidPrice)", decimal_or_none(result.order_summary, "gross_order_amount")),
        ("รายได้เข้า Wallet (finance_income)", expenses.get("order_income")),
        ("ค่าขนส่ง", expenses.get("shipping_fee")),
        ("ค่าบริการ", expenses.get("service_fee")),
        ("ค่าธรรมเนียมอื่น", expenses.get("other_fee")),
        ("คืนเงิน", expenses.get("refund")),
        ("รายการปรับปรุง", expenses.get("adjustment")),
        ("รายการไม่ทราบประเภท", expenses.get("unknown_amount")),
        ("ยอดสุทธิที่ควรเข้า Wallet", expenses.get("finance_net_amount")),
        ("ยอดโอนเข้าธนาคาร (ค่าบวกเพื่ออ่านง่าย)", wallet.get("bank_withdrawals")),
        ("ยอดถอน Auto Withdrawal (ค่าบวก)", wallet.get("auto_withdrawal_positive")),
        ("ยอดยกมา", wallet.get("opening_balance")),
        ("ยอดคงเหลือคำนวณ", wallet.get("calculated_closing_balance")),
        ("ยอดคงเหลือตามไฟล์", wallet.get("reported_closing_balance")),
        ("ผลต่าง Wallet", wallet.get("wallet_difference")),
        ("จำนวน MATCHED", matched),
        ("จำนวน NET_SETTLED", result.match_counts.get("NET_SETTLED", 0)),
        ("จำนวน ORDER_NOT_RELEASED", result.match_counts.get("ORDER_NOT_RELEASED", 0)),
        ("จำนวนรายการผิดปกติ (Exceptions)", exception_count),
        ("tolerance (บาท)", result.tolerance),
        ("สถานะรวม", result.overall_status),
    ]
    _write_kv(ws, 4, items, money_labels)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 18
    note_row = 4 + len(items) + 1
    ws.cell(row=note_row, column=1, value="หมายเหตุ").font = LABEL_FONT
    ws.cell(
        row=note_row + 1,
        column=1,
        value=(
            "จำนวนเงินในรายงานนี้คำนวณจาก signed amount ตามไฟล์ต้นทาง "
            "ค่าสัมบูรณ์ใช้เฉพาะตอนนำเสนอรายการถอนเข้าธนาคาร ค่าในวงเล็บหรือเครื่องหมายลบไม่ได้ถูกพลิกก่อนคำนวณ"
        ),
    )
    if result.warnings:
        ws.cell(row=note_row + 3, column=1, value="คำเตือน").font = WARN_FONT
        for i, warning in enumerate(result.warnings[:15]):
            ws.cell(row=note_row + 4 + i, column=1, value=warning)


def decimal_or_none(frame, column: str):
    from .normalizers import decimal_sum

    if frame is None or getattr(frame, "empty", True) or column not in frame.columns:
        return Decimal("0.00")
    return decimal_sum(frame[column])


def _write_order_summary(ws, result: ReconciliationResult) -> None:
    headers = [
        "order_number",
        "gross_order_amount",
        "order_status",
        "all_statuses",
        "mixed_status",
        "status_group",
        "source_line_count",
        "first_source_row",
        "last_source_row",
        "source_file",
    ]
    records = _records_from_frame(result.order_summary)
    last = _write_table(ws, headers, records)
    _add_generated_total(ws, headers, records, {"gross_order_amount"}, last)
    _autosize(ws)


def _write_finance_summary(ws, result: ReconciliationResult) -> None:
    headers = [
        "order_number",
        "finance_income",
        "shipping_fee",
        "service_fee",
        "other_fee",
        "refund",
        "adjustment",
        "unknown_amount",
        "finance_net_amount",
        "source_line_count",
        "first_source_row",
        "last_source_row",
        "source_file",
        "has_unknown_mapping",
    ]
    records = _records_from_frame(result.finance_summary)
    last = _write_table(ws, headers, records)
    _add_generated_total(
        ws,
        headers,
        records,
        {
            "finance_income",
            "shipping_fee",
            "service_fee",
            "other_fee",
            "refund",
            "adjustment",
            "unknown_amount",
            "finance_net_amount",
        },
        last,
    )
    _autosize(ws)


def _write_order_vs_finance(ws, result: ReconciliationResult) -> None:
    headers = [
        "order_number",
        "gross_order_amount",
        "finance_income",
        "finance_net_amount",
        "order_to_finance_difference",
        "implied_fee_or_net_gap",
        "order_status",
        "mixed_status",
        "match_status",
        "possible_reason",
        "in_orders",
        "in_finance",
        "order_source_file",
        "order_source_row_first",
        "finance_source_file",
        "finance_source_row_first",
    ]
    records = _records_from_frame(result.order_vs_finance)
    last = _write_table(ws, headers, records)
    _add_generated_total(
        ws,
        headers,
        records,
        {
            "gross_order_amount",
            "finance_income",
            "finance_net_amount",
            "order_to_finance_difference",
            "implied_fee_or_net_gap",
        },
        last,
    )
    if records:
        diff_col = get_column_letter(headers.index("order_to_finance_difference") + 1)
        status_col = get_column_letter(headers.index("match_status") + 1)
        ws.conditional_formatting.add(
            f"{diff_col}2:{diff_col}{last}",
            FormulaRule(
                formula=[f"ABS({diff_col}2)>{float(result.tolerance)}"],
                fill=FAIL_FILL,
            ),
        )
        ws.conditional_formatting.add(
            f"{status_col}2:{status_col}{last}",
            CellIsRule(operator="equal", formula=['"MATCHED"'], fill=PASS_FILL),
        )
        ws.conditional_formatting.add(
            f"{status_col}2:{status_col}{last}",
            CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL),
        )
    _autosize(ws)
    # Keep GENERATED TOTAL outside the autofilter range conceptually; still visible.


def _write_fee_details(ws, result: ReconciliationResult) -> None:
    headers = [
        "order_number",
        "freename",
        "mapped_category",
        "signed_amount",
        "absolute_amount",
        "mapping_known",
        "finance_profile",
        "transfer_status",
        "statement_code",
        "source_file",
        "source_row",
    ]
    records = _records_from_frame(result.finance_details)
    last = _write_table(ws, headers, records)
    _add_generated_total(ws, headers, records, {"signed_amount", "absolute_amount"}, last)
    _autosize(ws)


def _write_wallet(ws, result: ReconciliationResult) -> None:
    ws["A1"] = "สรุป Wallet / โอนเข้าธนาคาร"
    ws["A1"].font = TITLE_FONT
    wt = result.wallet_totals
    items = [
        ("opening_balance", wt.get("opening_balance")),
        ("wallet_inflows (Settlement signed)", wt.get("wallet_inflows")),
        ("wallet_adjustments", wt.get("wallet_adjustments")),
        ("withdrawals_signed", wt.get("withdrawals_signed")),
        ("bank_withdrawals (แสดงเป็นค่าบวก)", wt.get("bank_withdrawals")),
        ("auto_withdrawal_signed", wt.get("auto_withdrawal_signed")),
        ("auto_withdrawal_positive", wt.get("auto_withdrawal_positive")),
        ("signed_net", wt.get("signed_net")),
        ("calculated_closing_balance", wt.get("calculated_closing_balance")),
        ("reported_closing_balance", wt.get("reported_closing_balance")),
        ("wallet_difference", wt.get("wallet_difference")),
        ("formula", wt.get("formula")),
    ]
    _write_kv(ws, 3, items)
    headers = [
        "transaction_number",
        "transaction_time",
        "wallet_type",
        "wallet_sub_type",
        "signed_amount",
        "absolute_amount",
        "mapped_category",
        "is_bank_transfer",
        "is_auto_withdrawal",
        "remarks",
        "source_file",
        "source_row",
    ]
    start = 3 + len(items) + 2
    records = _records_from_frame(result.wallet_details)
    for c_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start, column=c_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    money_cols = {"signed_amount", "absolute_amount"}
    for r_idx, record in enumerate(records, start=start + 1):
        for c_idx, header in enumerate(headers, start=1):
            value = _excel_value(record.get(header))
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = THIN
            cell.font = NORMAL_FONT
            if header in money_cols and isinstance(value, (int, float)):
                cell.number_format = MONEY_FORMAT
    last = start + len(records)
    if records:
        ws.freeze_panes = f"A{start + 1}"
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{last}"
        _add_generated_total_at(ws, headers, money_cols, start, last)
    _autosize(ws)


def _add_generated_total_at(ws, headers, money_cols, header_row, last_data_row) -> None:
    total_row = last_data_row + 1
    for c_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=total_row, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.font = LABEL_FONT
        cell.border = THIN
        if c_idx == 1:
            cell.value = "GENERATED TOTAL"
            continue
        if header in money_cols:
            letter = get_column_letter(c_idx)
            cell.value = f"=SUM({letter}{header_row + 1}:{letter}{last_data_row})"
            cell.number_format = MONEY_FORMAT


def _write_exceptions(ws, result: ReconciliationResult) -> None:
    headers = [
        "exception_type",
        "severity",
        "orderNumber",
        "description",
        "expected_amount",
        "actual_amount",
        "difference",
        "possible_reason",
        "source_file",
        "source_row",
    ]
    records = _records_from_frame(result.exceptions)
    _write_table(ws, headers, records)
    if records:
        sev_col = get_column_letter(headers.index("severity") + 1)
        last = 1 + len(records)
        ws.conditional_formatting.add(
            f"{sev_col}2:{sev_col}{last}",
            CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL),
        )
        ws.conditional_formatting.add(
            f"{sev_col}2:{sev_col}{last}",
            CellIsRule(operator="equal", formula=['"WARNING"'], fill=WARN_FILL),
        )
    _autosize(ws)


def _write_unknown(ws, result: ReconciliationResult) -> None:
    headers = ["source", "raw_value", "field", "row_count", "signed_amount_sum", "suggested_mapping"]
    records = _records_from_frame(result.unknown_mappings)
    last = _write_table(ws, headers, records)
    _add_generated_total(ws, headers, records, {"signed_amount_sum"}, last)
    _autosize(ws)


def _write_validations(ws, result: ReconciliationResult) -> None:
    headers = ["validation_name", "status", "expected", "actual", "difference", "explanation"]
    records = _records_from_frame(result.validations)
    _write_table(ws, headers, records)
    _autosize(ws)


def _write_methodology(ws, result: ReconciliationResult) -> None:
    ws["A1"] = "Methodology / สูตรและสมมติฐาน"
    ws["A1"].font = TITLE_FONT
    ws.freeze_panes = "A2"
    lines = [
        f"โปรแกรม: {PROGRAM_NAME} {PROGRAM_VERSION}",
        f"เดือน: {result.month}",
        f"tolerance: {result.tolerance} บาท",
        "",
        "สูตร Order",
        "  gross_order_amount = sum(paidPrice) group by orderNumber",
        "  ตัดแถว Grand Total / Total ออกก่อนรวม",
        "  เลขออเดอร์เก็บเป็น string",
        "",
        "สูตร Finance",
        "  finance_net_amount = sum(signed_amount) ของทุกรายการในออเดอร์ (ไม่ใช้ค่าสัมบูรณ์)",
        "  finance_income = sum(signed_amount ที่ mapped เป็น order_income)",
        "  order_to_finance_difference = gross_order_amount - finance_income",
        "  หมวดหมู่: order_income, shipping_fee, service_fee, other_fee, refund, adjustment, unknown",
        "  Freename เท่ากับ 'Shipping fee Voucher Refund to Lazada' = shipping_fee",
        "  ค่าที่ไม่มี mapping = unknown และแสดงใน Unknown_Mappings",
        "",
        "สูตร Wallet",
        "  ใช้คอลัมน์ Amount เท่านั้น ห้ามใช้ Sub Type เป็นจำนวนเงิน",
        "  รายการ Type=Withdrawal / Sub Type=Auto Withdrawal = โอนเข้าธนาคาร",
        "  หาก Amount เป็นค่าติดลบ: signed_net = sum(signed_amount)",
        "  calculated_closing_balance = opening_balance + signed_net",
        "  เทียบเท่า opening + inflows + adjustments - withdrawals_as_positive",
        "  ถ้าไม่มี opening หรือ reported closing balance ให้แสดง 'ข้อมูลไม่เพียงพอ' ห้ามสมมติเป็น 0",
        "  รายงานยอดโอนเข้าธนาคารแสดงเป็นค่าบวกเพื่ออ่านง่าย แต่สูตรใช้ signed amount",
        "",
        "match_status",
        "  MATCHED = มีทั้ง Order และ Finance และผลต่างอยู่ใน tolerance",
        "  NET_SETTLED = มีทั้งสองไฟล์ แต่ Finance เป็นยอดสุทธิ (Income Order Overview) จึงไม่เท่า paidPrice",
        "  ORDER_NOT_RELEASED = มีใน Order แต่ยังไม่มีใน Finance",
        "  FINANCE_FROM_OTHER_PERIOD = มีใน Finance แต่ไม่มีใน Order",
        "  AMOUNT_MISMATCH = มีทั้งสองไฟล์ ยอดไม่ตรง และไม่ใช่ไฟล์ยอดสุทธิ",
        "  CANCELLED_OR_REFUNDED = สถานะยกเลิกหรือคืนเงิน",
        "  UNKNOWN = วิเคราะห์ไม่ได้",
        "",
        "การ join ใช้ full outer join ที่ orderNumber ทั้งสองทิศ ไม่ใช้ inner join และไม่ใช้ XLOOKUP ทางเดียว",
        "แถว GENERATED TOTAL สร้างจากสูตร Excel บนข้อมูลที่เขียนแล้ว ห้ามนำกลับไปคำนวณซ้ำในโปรแกรม",
        "",
        "สมมติฐานและข้อจำกัด",
    ]
    lines.extend(f"  - {note}" for note in result.methodology_notes)
    if result.warnings:
        lines.append("")
        lines.append("Warnings")
        lines.extend(f"  - {warning}" for warning in result.warnings)
    lines.extend(
        [
            "",
            "โปรแกรมไม่ใช่คำรับรองทางบัญชี นักบัญชีต้องตรวจผลอีกครั้งก่อนลงบัญชี",
            "ไฟล์ต้นทางใน input/ ไม่ถูกแก้ไข รายงานนี้อ้างอิง source_file และ source_row เพื่อตรวจสอบย้อนกลับ",
        ]
    )
    for idx, line in enumerate(lines, start=3):
        ws.cell(row=idx, column=1, value=line).font = NORMAL_FONT
    ws.column_dimensions["A"].width = 140
