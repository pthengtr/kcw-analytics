"""Shared Excel builders for unit tests. Synthetic data only — no real customers."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def write_sheet(
    path: Path,
    headers: list[str],
    rows: list[list[object]],
    *,
    title_rows: list[list[object]] | None = None,
    sheet_name: str = "Sheet1",
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    r = 1
    if title_rows:
        for title in title_rows:
            for c, value in enumerate(title, start=1):
                ws.cell(row=r, column=c, value=value)
            r += 1
    for c, header in enumerate(headers, start=1):
        ws.cell(row=r, column=c, value=header)
    r += 1
    for row in rows:
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
        r += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_orders(path: Path, rows: list[list[object]], title_rows=None) -> Path:
    headers = ["orderItemId", "orderNumber", "paidPrice", "status", "createTime"]
    return write_sheet(path, headers, rows, title_rows=title_rows, sheet_name="sheet1")


def write_finance_overview(path: Path, rows: list[list[object]], title_rows=None) -> Path:
    headers = [
        "หมายเลขคำสั่งซื้อ",
        "วันที่สร้างคำสั่งซื้อ",
        "รหัสสินค้าในคำสั่งซื้อ",
        "ชื่อสินค้า",
        "จำนวนเงิน(รวมภาษี)",
        "VAT Amount",
        "สถานะคำสั่งซื้อ",
        "ระยะเวลาใบแจ้งยอด",
        "รหัสรอบบิล",
        "สถานะการโอนเงิน",
        "Short Code",
    ]
    return write_sheet(path, headers, rows, title_rows=title_rows, sheet_name="Income Order Overview")


def write_finance_txn(path: Path, rows: list[list[object]], title_rows=None) -> Path:
    headers = [
        "หมายเลขคำสั่งซื้อ",
        "ชื่อรายการธุรกรรม",
        "จำนวนเงิน(รวมภาษี)",
        "VAT Amount",
        "สถานะการโอนเงิน",
        "สถานะคำสั่งซื้อ",
    ]
    return write_sheet(path, headers, rows, title_rows=title_rows, sheet_name="Income Overview")


def write_wallet(path: Path, rows: list[list[object]], title_rows=None) -> Path:
    headers = ["Transaction Number", "Transaction Time", "Type", "Sub Type", "Amount", "Remarks"]
    return write_sheet(path, headers, rows, title_rows=title_rows, sheet_name="Balance Transactions")
