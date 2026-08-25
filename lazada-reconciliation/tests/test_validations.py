"""Validation rules and report integrity."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.loaders import load_finance, load_orders, load_wallet
from src.main import infer_month, run
from src.reconciliation import run_reconciliation
from src.report_writer import write_report
from src.validations import run_validations
from tests.helpers import write_finance_overview, write_finance_txn, write_orders, write_wallet


def _prepare(tmp_path: Path):
    orders_dir = tmp_path / "orders"
    finance_dir = tmp_path / "finance"
    wallet_dir = tmp_path / "wallet"
    out_dir = tmp_path / "output"
    orders_dir.mkdir()
    finance_dir.mkdir()
    wallet_dir.mkdir()
    write_orders(
        orders_dir / "orders.xlsx",
        [
            ["I1", "1001", "100.00", "confirmed", "01 Aug 2026"],
            ["I2", "1001", "50.00", "confirmed", "01 Aug 2026"],
            ["I3", "1002", "80.00", "canceled", "02 Aug 2026"],
        ],
    )
    write_finance_overview(
        finance_dir / "finance.xlsx",
        [
            ["1001", "01 Aug 2026", "I1", "SKU", "120.00", "0", "ยืนยันแล้ว", "p", "B1", "โอนเงินไปยังยอดของฉันแล้ว", "SC"],
            ["", "01 Aug 2026", "", "", "-10.00", "0", "", "p", "B1", "โอนเงินไปยังยอดของฉันแล้ว", "SC"],
        ],
    )
    write_wallet(
        wallet_dir / "wallet.xlsx",
        [
            ["T1", "01 Aug 2026", "Deposit", "Settlement", "+120.00", "Statement"],
            ["T2", "02 Aug 2026", "Withdrawal", "Auto Withdrawal", "-90.00", "Paid"],
        ],
    )
    return orders_dir, finance_dir, wallet_dir, out_dir


def test_pre_post_group_totals_and_unknown_blank_order(tmp_path: Path):
    orders_dir, finance_dir, wallet_dir, _ = _prepare(tmp_path)
    orders = load_orders(orders_dir)
    finance = load_finance(finance_dir)
    wallet = load_wallet(wallet_dir)
    result = run_reconciliation(
        orders, finance, wallet, month="2026-08", generated_at="t", tolerance=Decimal("0.01")
    )
    result = run_validations(result, orders, finance, wallet)
    by_name = result.validations.set_index("validation_name")
    assert by_name.loc["pre_post_group_total_orders", "status"] == "PASS"
    assert by_name.loc["pre_post_group_total_finance", "status"] == "PASS"
    assert by_name.loc["opening_balance_present", "status"] == "WARNING"
    assert by_name.loc["empty_order_number", "status"] == "WARNING"
    assert result.overall_status in {"WARNING", "FAIL"}
    assert "กระทบยอดสำเร็จ" not in str(result.warnings)


def test_finance_components_equal_net(tmp_path: Path):
    orders_dir = tmp_path / "o"
    finance_dir = tmp_path / "f"
    wallet_dir = tmp_path / "w"
    orders_dir.mkdir()
    finance_dir.mkdir()
    wallet_dir.mkdir()
    write_orders(orders_dir / "o.xlsx", [["I1", "1001", "200.00", "confirmed", "01 Aug 2026"]])
    write_finance_txn(
        finance_dir / "f.xlsx",
        [
            ["1001", "Item price", "200.00", "0", "โอน", "ยืนยันแล้ว"],
            ["1001", "Commission", "-20.00", "0", "โอน", "ยืนยันแล้ว"],
            ["1001", "Shipping fee Voucher Refund to Lazada", "-5.00", "0", "โอน", "ยืนยันแล้ว"],
        ],
    )
    write_wallet(wallet_dir / "w.xlsx", [["T1", "t", "Deposit", "Settlement", "+175.00", "s"]])
    orders = load_orders(orders_dir)
    finance = load_finance(finance_dir)
    wallet = load_wallet(wallet_dir)
    result = run_reconciliation(
        orders, finance, wallet, month="2026-08", generated_at="t", tolerance=Decimal("0.01")
    )
    result = run_validations(result, orders, finance, wallet)
    assert result.expense_summary["finance_net_amount"] == Decimal("175.00")
    assert result.expense_summary["component_sum"] == Decimal("175.00")
    by_name = result.validations.set_index("validation_name")
    assert by_name.loc["finance_net_vs_components", "status"] == "PASS"


def test_report_has_all_sheets_and_generated_total(tmp_path: Path):
    orders_dir, finance_dir, wallet_dir, out_dir = _prepare(tmp_path)
    code = run(
        [
            "--orders",
            str(orders_dir),
            "--finance",
            str(finance_dir),
            "--wallet",
            str(wallet_dir),
            "--output",
            str(out_dir),
            "--month",
            "2026-08",
        ]
    )
    assert code in {0, 2}
    report = out_dir / "lazada_reconciliation_2026-08.xlsx"
    assert report.exists()
    wb = load_workbook(report)
    expected = {
        "Executive_Summary",
        "Order_Summary",
        "Finance_Summary",
        "Order_vs_Finance",
        "Fee_Details",
        "Wallet_Summary",
        "Exceptions",
        "Unknown_Mappings",
        "Validations",
        "Methodology",
    }
    assert expected <= set(wb.sheetnames)
    order_sheet = wb["Order_Summary"]
    values = [row[0].value for row in order_sheet.iter_rows(min_col=1, max_col=1)]
    assert "GENERATED TOTAL" in values
    # Must not merge cells in the data table.
    assert order_sheet.merged_cells.ranges == set() or len(order_sheet.merged_cells.ranges) == 0
    exec_sheet = wb["Executive_Summary"]
    assert exec_sheet["A2"].value.startswith("สถานะรวม")
    assert "กระทบยอดสำเร็จ" not in "".join(
        str(c.value or "") for row in exec_sheet.iter_rows(max_row=8) for c in row
    )


def test_infer_month_from_filename():
    month, warnings = infer_month(None, ["7LAZ1_orders_aug_2026.xlsx"])
    assert month == "2026-08"
    month2, warnings2 = infer_month(None, ["unknown.xlsx"])
    assert warnings2
    assert month2 != "2026-08"
