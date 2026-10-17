from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from src.kcw.online_statement_link import (
    BankCredit,
    Payout,
    assign_tiktok,
    collect_payouts,
    link_bills,
    match_bank_deposits,
    match_order,
    parse_lazada_rows,
    parse_shopee_income_rows,
    parse_shopee_rows,
    parse_tiktok_rows,
    reconcile_shopee_pickup_fees,
    staging_workbook_path,
    upload_destination,
    workbook_parse_error,
)

BKK = ZoneInfo("Asia/Bangkok")


def _dt(text: str) -> datetime:
    return datetime.strptime(text, "%Y-%m-%d %H:%M:%S").replace(tzinfo=BKK)


def test_shopee_withdrawal_collects_orders_and_fees():
    rows = [
        ("รายงาน",),
        ("วันที่", "ประเภทการทำธุรกรรม", "คำอธิบาย", "รหัสคำสั่งซื้อ", "รูปแบบธุรกรรม", "จำนวนเงิน", "สถานะ", "ยอด"),
        ("2026-09-15 01:45:10", "การถอนเงิน", "การถอนเงินอัตโนมัติ", "-", "เงินออก", -300, "ทำรายการสำเร็จ", 0),
        ("2026-09-14 12:00:00", "รายรับจากคำสั่งซื้อ", "เงินโอนจากคำสั่งซื้อ #260909J3DR9V65", "260909J3DR9V65", "เงินเข้า", 200, "ทำรายการสำเร็จ", 300),
        ("2026-09-14 10:00:00", "รายการปรับปรุง", "Pick up fee", "-", "เงินออก", -50, "ทำรายการสำเร็จ", 100),
        ("2026-09-13 10:00:00", "รายรับจากคำสั่งซื้อ", "เงินโอนจากคำสั่งซื้อ #260908FKU2X0AC", "260908FKU2X0AC", "เงินเข้า", 150, "ทำรายการสำเร็จ", 150),
        ("2026-09-16 09:00:00", "รายรับจากคำสั่งซื้อ", "เงินโอนจากคำสั่งซื้อ #260910AAAAAA11", "260910AAAAAA11", "เงินเข้า", 40, "ทำรายการสำเร็จ", 40),
    ]
    payouts = parse_shopee_rows(rows, shop="SP", source_file="Shopee/SP/a.xlsx")
    transferred = [p for p in payouts if p.status == "transferred"]
    pending = [p for p in payouts if p.status == "pending"]
    assert len(transferred) == 1
    assert transferred[0].amount == Decimal("300.00")
    assert {line.order_id for line in transferred[0].lines if line.order_id} == {
        "260909J3DR9V65",
        "260908FKU2X0AC",
    }
    fees = [line for line in transferred[0].lines if line.kind == "adjustment"]
    assert fees[0].net == Decimal("-50.00")
    assert sum((line.net for line in transferred[0].lines), Decimal("0")) == Decimal("300.00")
    assert pending[0].lines[0].order_id == "260910AAAAAA11"


def test_shopee_income_keeps_each_fee_and_adjustment():
    income = [
        ("ชื่อผู้ใช้ (ผู้ขาย)", "จาก", "ถึง"),
        ("kc.industry", "2026-09-14", "2026-09-25"),
        ("ยอดรวม (฿)",),
        (
            "ลำดับที่",
            "หมายเลขคำสั่งซื้อ",
            "ช่องทางการชำระเงินของผู้ซื้อ",
            "วันที่โอนชำระเงินสำเร็จ",
            "สินค้าราคาปกติ",
            "ค่าจัดส่งที่ชำระโดยผู้ซื้อ",
            "ค่าจัดส่งที่ Shopee ชำระโดยชื่อของคุณ",
            "ค่าคอมมิชชั่น",
            "ค่าธุรกรรมการชำระเงิน",
            "ค่าธรรมเนียมเติมเงินโฆษณาจากเงิน Escrow",
            "โบนัสส่วนลดเครื่องเก่าแลกใหม่จากผู้ขาย",
            "จำนวนเงินทั้งหมดที่โอนแล้ว (฿)",
        ),
        (
            1,
            "260911R83ACM85",
            "เก็บเงินปลายทาง",
            "2026-09-14",
            1350,
            38,
            -38,
            -159,
            -45,
            -15,
            0,
            1131,
        ),
    ]
    adjustments = [
        (
            "ลำดับที่",
            "วันที่โอนยอดเงินที่ปรับปรุง",
            "ประเภทการปรับปรุง | คำอธิบาย",
            "เหตุผลในการปรับปรุงยอดเงิน",
            "ยอดที่ปรับปรุง",
            "รหัสคำสั่งซื้อที่เกี่ยวข้อง",
        ),
        (1, "2026-09-19", "การคืนเงิน/คืนสินค้าหลังคำสั่งซื้อสำเร็จ", "", -360, "260908GSA520FV"),
    ]
    summary = [("จาก", "2026-09-14"), ("ถึง", "2026-09-25")]
    payouts = parse_shopee_income_rows(
        income,
        adjustments,
        summary,
        shop="SP",
        source_file="Shopee/SP/income.xlsx",
    )
    assert len(payouts) == 1
    payout = payouts[0]
    assert payout.reference == "2026-09-14 – 2026-09-25"
    assert payout.amount == Decimal("771.00")
    order = payout.lines[0]
    assert order.order_id == "260911R83ACM85"
    assert order.gross == Decimal("1388.00")
    assert order.expense == Decimal("-257.00")
    assert order.net == Decimal("1131.00")
    assert {fee.name: fee.amount for fee in order.fees if fee.amount < 0} == {
        "ค่าจัดส่งที่ Shopee ชำระโดยชื่อของคุณ": Decimal("-38.00"),
        "ค่าคอมมิชชั่น": Decimal("-159.00"),
        "ค่าธุรกรรมการชำระเงิน": Decimal("-45.00"),
        "ค่าธรรมเนียมเติมเงินโฆษณาจากเงิน Escrow": Decimal("-15.00"),
    }
    refund = payout.lines[1]
    assert refund.kind == "adjustment"
    assert refund.order_id == "260908GSA520FV"
    assert refund.net == Decimal("-360.00")


def test_shopee_pickup_fee_aligns_income_total_with_the_bank():
    payouts = parse_shopee_income_rows(
        [
            (
                "ลำดับที่",
                "หมายเลขคำสั่งซื้อ",
                "วันที่โอนชำระเงินสำเร็จ",
                "สินค้าราคาปกติ",
                "ค่าคอมมิชชั่น",
                "โบนัสส่วนลดเครื่องเก่าแลกใหม่จากผู้ขาย",
                "จำนวนเงินทั้งหมดที่โอนแล้ว (฿)",
            ),
            (1, "260911R83ACM85", "2026-09-14", 1350, -219, 0, 1131),
        ],
        [],
        [("จาก", "2026-09-14"), ("ถึง", "2026-09-25")],
        shop="IAPS",
        source_file="Shopee/IAPS/income.xlsx",
    )
    credits = [
        BankCredit("line-1", date(2026, 9, 25), Decimal("1101.00"), "TR from 9825080752 Shopeepay"),
        BankCredit("line-2", date(2026, 9, 25), Decimal("7021.00"), "TR from 9825080752 Shopeepay"),
    ]
    reconcile_shopee_pickup_fees(payouts, credits)
    assert payouts[0].amount == Decimal("1101.00")
    pickup = payouts[0].lines[-1]
    assert pickup.fee_name == "Charging Pick up fee"
    assert pickup.net == Decimal("-30.00")
    deposits = match_bank_deposits(payouts, credits)
    assert [item.statement_line_id for item in deposits] == ["line-1"]
    assert deposits[0].shop == "IAPS"
    assert deposits[0].amount == Decimal("1101.00")
    links = link_bills(
        payouts,
        {
            "260911R83ACM85": [
                {
                    "billno": "TAD1",
                    "po": "260911R83ACM85",
                    "bill_date": "2026-09-14",
                    "aftertax": "1131",
                    "acctname": "shop",
                    "canceled": False,
                }
            ]
        },
    )
    assert links[0].payout_key == payouts[0].payout_key
    assert links[0].payout_key.endswith("|1101.00")


def test_lazada_statement_is_the_payout():
    header = [""] * 16
    header[1] = "Transaction Type"
    header[2] = "Fee Name"
    header[7] = "Amount"
    header[11] = "Statement"
    header[12] = "Paid Status"
    header[13] = "Order No."
    rows = [
        tuple(header),
        ("20-Sep-2026", "Orders-Sales", "Item Price Credit", "", "", "", "", 100, "", "", "", "20 Sep 2026 - 20 Sep 2026", "paid", "1118594800859773"),
        ("20-Sep-2026", "Orders-Lazada Fees", "Commission", "", "", "", "", -10, "", "", "", "20 Sep 2026 - 20 Sep 2026", "paid", "1118594800859773"),
        ("19-Sep-2026", "Orders-Sales", "Item Price Credit", "", "", "", "", 50, "", "", "", "19 Sep 2026 - 19 Sep 2026", "paid", "1111111111111111"),
    ]
    payouts = parse_lazada_rows(rows, shop="LAZ1", source_file="Lazada/LAZ1/a.xlsx")
    by_ref = {p.reference: p for p in payouts}
    day = by_ref["20 Sep 2026 - 20 Sep 2026"]
    assert day.amount == Decimal("90.00")
    assert day.status == "transferred"
    assert day.lines[0].order_id == "1118594800859773"
    assert day.lines[0].expense == Decimal("-10.00")
    assert len(day.lines[0].fees) == 2


def test_tiktok_orders_follow_the_withdrawal_that_sweeps_them():
    order_header = [""] * 16
    order_header[0] = "หมายเลขคำสั่งซื้อ/การปรับ"
    order_header[1] = "ประเภทธุรกรรม"
    order_header[3] = "เวลาที่ชำระคำสั่งซื้อ"
    order_header[5] = "ยอดการชำระเงินทั้งหมด"
    order_header[6] = "รายได้ทั้งหมด"
    order_header[13] = "ค่าธรรมเนียมทั้งหมด"
    order_header[14] = "ค่าคอมมิชชั่น TikTok Shop"
    orders = [
        tuple(order_header),
        ("585982682407208429", "คำสั่งซื้อ", "", "2026/09/09", "", 80, 100, "", "", "", "", "", "", -20, -10),
        ("585982682407208430", "คำสั่งซื้อ", "", "2026/09/15", "", 40, 50, "", "", "", "", "", "", -10, -5),
        ("585982682407208431", "คำสั่งซื้อ", "", "2026/09/16", "", 15, 20, "", "", "", "", "", "", -5, -2),
    ]
    withdraw_header = ["ประเภทธุรกรรม", "ID อ้างอิง", "เวลาส่งคำขอ", "จำนวน", "สถานะ", "เวลาที่สำเร็จ", "บัญชีธนาคาร"]
    withdrawals = [
        tuple(withdraw_header),
        ("Withdrawal", "3701", "2026/09/09", -200, "Transferred", "2026/09/09", "********1139"),
        ("Withdrawal", "3703", "2026/09/16", -120, "Transferred", "2026/09/16", "********1139"),
        ("Earnings", "3702", "2026/09/15", 40, "Transferred", "2026/09/15", "/"),
    ]
    order_lines, payouts = parse_tiktok_rows(orders, withdrawals, shop="ICE", source_file="Tiktok/ICE/a.xlsx")
    assigned = assign_tiktok(order_lines, payouts, shop="ICE", source_file="Tiktok/ICE/a.xlsx")
    by_ref = {p.reference: p for p in assigned if p.status == "transferred"}
    pending = [p for p in assigned if p.status == "pending"]
    assert by_ref["3701"].amount == Decimal("200.00")
    assert by_ref["3701"].lines == []
    swept = {line.order_id for line in by_ref["3703"].lines}
    assert swept == {"585982682407208429", "585982682407208430"}
    assert by_ref["3703"].lines[0].fees[0].name == "ค่าคอมมิชชั่น TikTok Shop"
    assert pending[0].lines[0].order_id == "585982682407208431"


def test_po_match_prefers_exact_then_head_trim():
    index = {
        "260909J3DR9V65": [{"billno": "TAD1"}],
        "18594800859773": [{"billno": "TAD2"}],
        "982682407208429": [{"billno": "TAD3"}],
    }
    assert match_order("260909J3DR9V65", index)[0] == "exact"
    assert match_order("1118594800859773", index)[0] == "drop_front_2"
    assert match_order("585982682407208429", index)[0] == "drop_front_3"
    assert match_order("9999999999999999", index) is None


def test_collect_reads_three_platforms(tmp_path):
    laz = tmp_path / "Lazada" / "LAZ1"
    shopee = tmp_path / "Shopee" / "SP"
    tiktok = tmp_path / "Tiktok" / "ICE"
    for folder in (laz, shopee, tiktok):
        folder.mkdir(parents=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction Overview"
    ws.append(["Transaction Date", "Transaction Type", "Fee Name", "Transaction Number", "Details", "Seller SKU", "Lazada SKU", "Amount", "VAT in Amount", "WHT Amount", "WHT included in Amount", "Statement", "Paid Status", "Order No."])
    ws.append(["20-Sep-2026", "Orders-Sales", "Item Price Credit", "1", "", "", "", 10, 0, 0, "No", "20 Sep 2026 - 20 Sep 2026", "paid", "1118594800859773"])
    wb.save(laz / "a.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction Report"
    ws.append(["รายงาน"])
    ws.append(["วันที่", "ประเภทการทำธุรกรรม", "คำอธิบาย", "รหัสคำสั่งซื้อ", "รูปแบบธุรกรรม", "จำนวนเงิน", "สถานะ", "ยอด"])
    ws.append(["2026-09-15 01:45:10", "การถอนเงิน", "การถอนเงินอัตโนมัติ", "-", "เงินออก", -10, "ทำรายการสำเร็จ", 0])
    ws.append(["2026-09-14 12:00:00", "รายรับจากคำสั่งซื้อ", "x", "260909J3DR9V65", "เงินเข้า", 10, "ทำรายการสำเร็จ", 10])
    wb.save(shopee / "a.xlsx")

    wb = Workbook()
    orders = wb.active
    orders.title = "รายละเอียดคำสั่งซื้อ"
    orders.append(["หมายเลขคำสั่งซื้อ/การปรับ", "ประเภทธุรกรรม", "เวลาที่สร้างคำสั่งซื้อ", "เวลาที่ชำระคำสั่งซื้อ", "สกุลเงิน", "ยอดการชำระเงินทั้งหมด", "รายได้ทั้งหมด", "", "", "", "", "", "", "ค่าธรรมเนียมทั้งหมด"])
    orders.append(["585982682407208429", "คำสั่งซื้อ", "2026/09/09", "2026/09/09", "THB", 80, 100, "", "", "", "", "", "", -20])
    withdraw = wb.create_sheet("บันทึกการถอน")
    withdraw.append(["ประเภทธุรกรรม", "ID อ้างอิง", "เวลาส่งคำขอ", "จำนวน", "สถานะ", "เวลาที่สำเร็จ", "บัญชีธนาคาร"])
    withdraw.append(["Withdrawal", "3701", "2026/09/09", -80, "Transferred", "2026/09/09", "********1139"])
    wb.save(tiktok / "a.xlsx")

    payouts = collect_payouts(tmp_path)
    platforms = {p.platform for p in payouts if p.status == "transferred"}
    assert platforms == {"lazada", "shopee", "tiktok"}


def _day_payout(shop: str, day: date, amount: str) -> Payout:
    return Payout(
        platform="lazada",
        shop=shop,
        payout_at=datetime(day.year, day.month, day.day, tzinfo=BKK),
        amount=Decimal(amount),
        reference=day.isoformat(),
        status="transferred",
        source_file="Lazada/LAZ1/week.xlsx",
        note="",
    )


def test_bank_credit_links_the_statement_week_that_sums_to_it():
    days = [
        (date(2026, 9, 14), "3781.64"),
        (date(2026, 9, 15), "1746.76"),
        (date(2026, 9, 16), "3032.49"),
        (date(2026, 9, 17), "1191.58"),
        (date(2026, 9, 18), "3701.51"),
        (date(2026, 9, 19), "2838.19"),
        (date(2026, 9, 20), "7198.23"),
    ]
    payouts = [_day_payout("LAZ1", day, amount) for day, amount in days]
    payouts.append(_day_payout("LPNT", date(2026, 9, 20), "4064.55"))
    deposits = match_bank_deposits(
        payouts,
        [
            BankCredit("line-1", date(2026, 9, 23), Decimal("23490.40"), "BPS/017/01/Lazada Ltd./108682"),
            BankCredit("line-2", date(2026, 9, 23), Decimal("4064.55"), "BPS/017/01/Lazada Ltd./108682"),
        ],
    )
    by_line = {item.statement_line_id: item for item in deposits}
    week = by_line["line-1"]
    assert week.shop == "LAZ1"
    assert week.period_from == date(2026, 9, 14)
    assert week.period_to == date(2026, 9, 20)
    assert len(week.payout_keys) == 7
    assert by_line["line-2"].shop == "LPNT"


def test_bank_credit_skips_an_amount_shared_by_two_shops():
    payouts = [
        _day_payout("LAZ1", date(2026, 9, 20), "100.00"),
        _day_payout("LPNT", date(2026, 9, 20), "100.00"),
    ]
    deposits = match_bank_deposits(
        payouts,
        [BankCredit("line-1", date(2026, 9, 23), Decimal("100.00"), "Lazada Ltd.")],
    )
    assert deposits == []


def _save(path, sheets: dict[str, list[tuple]]) -> None:
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for row in rows:
            ws.append(row)
    wb.save(path)


def test_upload_destination_uses_the_drive_shop_folder(tmp_path):
    assert upload_destination(tmp_path, "lazada", "LAZ1", "LAZ1 week.xlsx") == tmp_path / "Lazada" / "LAZ1" / "LAZ1 week.xlsx"


def test_staging_workbook_keeps_an_xlsx_suffix(tmp_path):
    dest = upload_destination(tmp_path, "lazada", "LICE", "LICE_week_28-04.xlsx")
    staging = staging_workbook_path(dest)
    assert staging == dest.with_name(".LICE_week_28-04.upload.xlsx")
    assert staging.suffix == ".xlsx"
    dest.parent.mkdir(parents=True)
    _save(staging, {
        "Transaction Overview": [
            ("Amount", "Statement", "Paid Status", "Order No.", "Fee Name", "Transaction Type"),
            (10, "20 Sep 2026 - 20 Sep 2026", "paid", "1118594800859773", "Item Price Credit", "Orders-Sales"),
        ]
    })
    assert workbook_parse_error(staging, "lazada", "LICE") is None


def test_workbook_parse_accepts_and_rejects_each_format(tmp_path):
    lazada = tmp_path / "laz.xlsx"
    _save(lazada, {
        "Transaction Overview": [
            ("Amount", "Statement", "Paid Status", "Order No.", "Fee Name", "Transaction Type"),
            (10, "20 Sep 2026 - 20 Sep 2026", "paid", "1118594800859773", "Item Price Credit", "Orders-Sales"),
        ]
    })
    assert workbook_parse_error(lazada, "lazada", "LAZ1") is None
    bad = tmp_path / "bad-laz.xlsx"
    _save(bad, {"Sheet1": [("hello",)]})
    assert workbook_parse_error(bad, "lazada", "LAZ1")

    shopee = tmp_path / "sp.xlsx"
    _save(shopee, {
        "Transaction Report": [
            ("รายงาน",),
            ("วันที่", "ประเภทการทำธุรกรรม", "คำอธิบาย", "รหัสคำสั่งซื้อ", "รูปแบบธุรกรรม", "จำนวนเงิน", "สถานะ", "ยอด"),
            ("2026-09-15 01:45:10", "การถอนเงิน", "การถอนเงินอัตโนมัติ", "-", "เงินออก", -10, "ทำรายการสำเร็จ", 0),
        ]
    })
    assert workbook_parse_error(shopee, "shopee", "SP") is None
    income = tmp_path / "sp-income.xlsx"
    _save(income, {
        "Summary": [("จาก", "2026-09-14"), ("ถึง", "2026-09-25")],
        "Income": [
            ("ลำดับที่", "หมายเลขคำสั่งซื้อ", "วันที่โอนชำระเงินสำเร็จ", "สินค้าราคาปกติ", "ค่าคอมมิชชั่น", "โบนัสส่วนลดเครื่องเก่าแลกใหม่จากผู้ขาย", "จำนวนเงินทั้งหมดที่โอนแล้ว (฿)"),
            (1, "260911R83ACM85", "2026-09-14", 1350, -219, 0, 1131),
        ],
    })
    assert workbook_parse_error(income, "shopee", "SP") is None
    bad_sp = tmp_path / "bad-sp.xlsx"
    _save(bad_sp, {"Transaction Report": [("วันที่", "อย่างอื่น")]})
    assert workbook_parse_error(bad_sp, "shopee", "SP")

    tiktok = tmp_path / "tt.xlsx"
    _save(tiktok, {
        "รายละเอียดคำสั่งซื้อ": [
            ("หมายเลขคำสั่งซื้อ/การปรับ", "ประเภทธุรกรรม", "เวลาที่ชำระคำสั่งซื้อ", "ยอดการชำระเงินทั้งหมด", "รายได้ทั้งหมด", "ค่าธรรมเนียมทั้งหมด"),
            ("585982682407208429", "คำสั่งซื้อ", "2026/09/09", 80, 100, -20),
        ],
        "บันทึกการถอน": [
            ("ประเภทธุรกรรม", "ID อ้างอิง", "เวลาส่งคำขอ", "จำนวน", "สถานะ", "เวลาที่สำเร็จ", "บัญชีธนาคาร"),
            ("Withdrawal", "3701", "2026/09/09", -80, "Transferred", "2026/09/09", "********1139"),
        ],
    })
    assert workbook_parse_error(tiktok, "tiktok", "ICE") is None
    bad_tt = tmp_path / "bad-tt.xlsx"
    _save(bad_tt, {"Sheet1": [("hello",)]})
    assert workbook_parse_error(bad_tt, "tiktok", "ICE")
